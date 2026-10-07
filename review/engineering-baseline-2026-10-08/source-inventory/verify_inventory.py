"""Independent read-only comparison against bundled openpyxl; never saves XLSX."""

import collections
import csv
import hashlib
import json
import pathlib
import sys

import openpyxl

OUT = pathlib.Path(__file__).resolve().parent
ROOT = OUT.parents[2]
DIM = json.loads((OUT / "dimensions-workbook.json").read_text(encoding="utf-8"))
REG = json.loads((OUT / "category-register-records.json").read_text(encoding="utf-8"))
SUMMARY = json.loads((OUT / "category-registers.json").read_text(encoding="utf-8"))
COMPARISON = json.loads((OUT / "snapshot-comparison.json").read_text(encoding="utf-8"))
FINGERPRINTS = json.loads((OUT / "source-fingerprints.json").read_text(encoding="utf-8"))
cell_checks = 0
formula_checks = 0
books_checked = 0


def check_book(relative, sheets):
    global cell_checks, formula_checks, books_checked
    cached = openpyxl.load_workbook(ROOT / relative, read_only=True, data_only=True)
    formula_book = openpyxl.load_workbook(ROOT / relative, read_only=True, data_only=False)
    try:
        for name, expected in sheets.items():
            actual = {c.coordinate: c.value for row in cached[name].iter_rows() for c in row if c.value is not None}
            actual_formulas = {c.coordinate: c.value for row in formula_book[name].iter_rows() for c in row if c.data_type == "f"}
            for address, item in expected.items():
                assert actual.get(address) == item["value"], (relative, name, address, actual.get(address), item["value"])
                cell_checks += 1
                if "formula" in item:
                    assert actual_formulas.get(address) == "=" + item["formula"], (relative, name, address, "formula mismatch")
                    formula_checks += 1
        books_checked += 1
    finally:
        cached.close()
        formula_book.close()


check_book(DIM["path"], {s["name"]: s["cells"] for s in DIM["sheets"]})
for register in REG.values():
    fields_by_sheet = {name: {field["cell"]: field for row in table["records"] for field in row["fields"].values()} for name, table in register["tables"].items()}
    check_book(register["path"], fields_by_sheet)

assert books_checked == 7
assert all(item["match"] for item in COMPARISON["comparisons"])
assert all(not ids for ids in SUMMARY["duplicates"].values())
assert all(not ids for ids in SUMMARY["retired"].values())
assert all(not ids for ids in SUMMARY["id_set_comparisons"].values())
assert not SUMMARY["nonempty_engineering_inputs"]
assert len(DIM["tables"]["Dimensions"]["records"]) == 463
assert len(DIM["tables"]["Model grid"]["records"]) == 51
assert len(DIM["tables"]["Open items"]["records"]) == 22
assert len(DIM["tables"]["Sources"]["records"]) == 9
assert len({row["id"] for row in DIM["tables"]["Dimensions"]["records"]}) == 463
assert len({row["id"] for row in DIM["tables"]["Open items"]["records"]}) == 22
csv_counts = {}
for name, expected_count in [("dimension-records.csv", 463), ("model-grid-records.csv", 51), ("open-items.csv", 22), ("drawing-sources.csv", 9), ("category-counts.csv", 6), ("engineering-input-columns.csv", 114), ("register-id-index.csv", 3337)]:
    with (OUT / name).open(encoding="utf-8", newline="") as handle:
        count = len(list(csv.DictReader(handle)))
    assert count == expected_count, (name, count, expected_count)
    csv_counts[name] = count
for entry in FINGERPRINTS["before"]:
    actual = hashlib.sha256((ROOT / entry["path"]).read_bytes()).hexdigest()
    assert actual == entry["sha256"], (entry["path"], "source fingerprint changed")
checks = {
    "result": "passed",
    "python_executable": sys.executable,
    "openpyxl_version": openpyxl.__version__,
    "workbook_mode": "read_only=True; data_only=True and data_only=False; no save operation",
    "workbooks_checked": books_checked,
    "cell_value_comparisons": cell_checks,
    "exact_formula_text_comparisons": formula_checks,
    "saved_fingerprint_and_count_comparisons": len(COMPARISON["comparisons"]),
    "csv_record_counts": csv_counts,
    "source_fingerprints_rechecked": len(FINGERPRINTS["before"]),
    "scope_limit": "Extraction accuracy and saved-file consistency only. No engineering adequacy or current native-Excel recalculation verified.",
}
(OUT / "verification.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(checks, ensure_ascii=False, indent=2))
