#!/usr/bin/env python3
"""Verify and package the dated LOCAL CAD exports. Nothing is uploaded."""
import hashlib
import json
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pymupdf as fitz

ROOT=Path(__file__).resolve().parent.parent
DEST=ROOT/'exports/dwg-review-2026-10-26'
data=json.loads((DEST/'register-data.json').read_text())
manifest=json.loads((DEST/'export-manifest.json').read_text())
wb=openpyxl.load_workbook(DEST/'04_Key_Dimensions.xlsx',data_only=True)
assert wb.sheetnames==['Key dimensions','Grid positions','Drawing index']
assert [list(r) for r in wb['Key dimensions'].iter_rows(min_row=7,values_only=True)]==data['keys']
assert [list(r) for r in wb['Grid positions'].iter_rows(min_row=7,values_only=True)]==data['grid']
for row in wb['Key dimensions'].iter_rows(min_row=7,values_only=True):
    assert abs(row[3]/1000-row[4])<1e-10
assert (DEST/'04_Key_Dimensions.xlsx').read_bytes()==(ROOT/'docs/layout_design/dwg-2026-10-23/Thach_Bi_Key_Dimensions.xlsx').read_bytes()
assert len(manifest['views'])==64
counts={'full_views':27,'detail_views':37,'key_witnesses':len(data['keys']),'grid_positions':len(data['grid'])}
for filename,expected in [('01_All_Drawing_Views.pdf',28),('02_Enlarged_Details.pdf',38),('03_Key_Dimensions_and_Guide.pdf',12)]:
    with fitz.open(DEST/filename) as doc:
        assert len(doc)==expected
        for page in doc:
            assert 'Not for construction' in page.get_text(), (filename,page.number)
            assert not page.get_images(), (filename,page.number,'unexpected raster image')
        counts[filename]={'pages':len(doc),'raster_images':0}
for view in manifest['views']:
    with fitz.open(DEST/view['file']) as doc:
        assert len(doc)==1 and not doc[0].get_images()
shutil.copy2(ROOT/'docs/layout_design/Thach_Bi_Church_Dimensions.xlsx',DEST/'05_Full_Dimension_Register.xlsx')
readme='''THACH BI CHURCH - LOCAL CAD REVIEW, 26 OCTOBER 2026

Start with 03_Key_Dimensions_and_Guide.pdf.
01_All_Drawing_Views.pdf: 27 full vector views, A2, plus index.
02_Enlarged_Details.pdf: 37 enlarged vector crops, A3, plus guide.
04_Key_Dimensions.xlsx: 79 key witnesses and 21 grid coordinates.
05_Full_Dimension_Register.xlsx: complete project register, older PDF evidence retained.
views/ and details/: one vector PDF per full view/crop.
export-manifest.json: exact source crop boundaries and source hash.

The source DWG saved 2026-10-23 is the owner-selected project design basis.
Its bytes match the already registered DWG revision. The original DWG is not modified.
Source tables are in source-tables/. DIMENSION IDs and handles link back to CAD.

IN DEVELOPMENT / DANG PHAT TRIEN - NOT FOR CONSTRUCTION / KHONG DUNG DE THI CONG
Vector linework stays sharp when zoomed. Page scales are fit-to-page, NOT calibrated.
Use stated dimensions/endpoints. Local u/v drawings do not supply global XYZ.
Native-CAD font/plot comparison, survey and responsible designer review remain unverified.
Internal platform, ridge, cornice and overall-extent conflicts remain explicitly open.
Designer drawings remain under publication hold H1. This local package is not permission
to publish or an approved fabrication, procurement or construction package.
'''
(DEST/'README.txt').write_text(readme)
counts['workbook_values_match_source']=True
counts['status']='Software/export verification only; internal conflicts and engineering holds retained.'
(DEST/'verification.json').write_text(json.dumps(counts,indent=2)+'\n')
files=[(DEST/name,name) for name in ['README.txt','01_All_Drawing_Views.pdf','02_Enlarged_Details.pdf',
        '03_Key_Dimensions_and_Guide.pdf','04_Key_Dimensions.xlsx','05_Full_Dimension_Register.xlsx',
        'export-manifest.json','source-audit.json','verification.json']]
files += [(DEST/v['file'],v['file']) for v in manifest['views']]
files += [(ROOT/'docs/layout_design/dwg-2026-10-23'/name,'source-tables/'+name)
          for name in ['dimensions.csv','levels.csv','grid-check.csv','facade.csv','key-dimensions.csv','sheets.json']]
zip_path=DEST/'Thach_Bi_CAD_Review.zip'
with zipfile.ZipFile(zip_path,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as archive:
    for file,name in files:
        archive.write(file,name)
with zipfile.ZipFile(zip_path) as archive:
    assert archive.testzip() is None
    assert len(archive.namelist())==len(files)
hashes={name:hashlib.sha256(file.read_bytes()).hexdigest() for file,name in files}
(DEST/'delivery-sha256.json').write_text(json.dumps(hashes,indent=2)+'\n')
print(json.dumps(counts,indent=2))
print(f'Local ZIP: {zip_path.relative_to(ROOT)}; {len(files)} files; {zip_path.stat().st_size/1e6:.1f} MB')
